"""
build_report_docx.py
Converts report/Project_Report.md into an editable Word document
(report/Project_Report.docx) with embedded figures and formatted tables.

Supports the Markdown subset used in the report: headings, paragraphs with
**bold** / *italic* / `code`, bullet and numbered lists, block quotes, pipe tables,
images and horizontal rules.

Run from the project root:  python report/build_report_docx.py
(or convert another Markdown file: python report/build_report_docx.py <input.md> <output.docx>)
"""

import os
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'Project_Report.md')
OUT = os.path.join(HERE, 'Project_Report.docx')

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)')


def add_inline(paragraph, text, bold=False, size=None):
    """Add text with **bold**, *italic* and `code` runs."""
    text = text.replace('\\*', '*')
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.bold = bold
        else:
            run = paragraph.add_run(part)
            run.bold = bold
        if size:
            run.font.size = size


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def split_row(row):
    """Split a Markdown table row on unescaped '|' and unescape '\\|' inside cells."""
    cells = re.split(r'(?<!\\)\|', row.strip().strip('|'))
    return [c.strip().replace('\\|', '|') for c in cells]


def add_table(doc, rows):
    header = split_row(rows[0].strip())
    body = [split_row(r) for r in rows[2:]]
    # table text: 10.5 pt for narrow tables, smaller only when many columns must fit the page
    size = Pt(10.5) if len(header) <= 5 else Pt(9.5) if len(header) <= 8 else Pt(8.5)
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ''
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        add_inline(cell.paragraphs[0], h, bold=True, size=size)
        shade(cell, 'D9E2F3')
    for i, r in enumerate(body, start=1):
        for j in range(len(header)):
            cell = table.rows[i].cells[j]
            cell.text = ''
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            add_inline(cell.paragraphs[0], r[j] if j < len(r) else '', size=size)
    doc.add_paragraph()


def main(src=SRC, out=OUT):
    base = os.path.dirname(os.path.abspath(src))
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(12)                       # Times New Roman 12 pt
    style.paragraph_format.line_spacing = 1.15     # 1.15 line spacing
    style.paragraph_format.space_after = Pt(6)
    for name in ('Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman') if st.element.rPr is not None else None
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.styles['Heading 2'].font.size = Pt(13)
    doc.styles['Heading 3'].font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)

    lines = open(src, encoding='utf-8').read().splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == '---':
            i += 1
            continue

        # Pipe table
        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        # Image
        m = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if m:
            path = os.path.normpath(os.path.join(base, m.group(2)))
            if os.path.exists(path):
                from PIL import Image as _Img
                with _Img.open(path) as im:
                    ratio = im.height / im.width
                # fit inside the printable area (15.5 cm wide, at most 21 cm tall)
                if 15.5 * ratio > 21:
                    doc.add_picture(path, height=Cm(21))
                else:
                    doc.add_picture(path, width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.paragraphs[-1].paragraph_format.keep_with_next = True   # keep figure with its caption
            else:
                doc.add_paragraph(f'[missing figure: {m.group(2)}]')
            i += 1
            continue

        # Headings
        m = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if first_heading:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(16)
                first_heading = False
            else:
                h = doc.add_heading(text, level=min(level - 1, 3) or 1)
                for r in h.runs:
                    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
                    r.font.name = 'Times New Roman'
            i += 1
            continue

        # Block quote
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_inline(p, stripped.lstrip('> ').strip())
            i += 1
            continue

        # Lists
        m_bullet = re.match(r'^[*\-]\s+(.*)', stripped)
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_bullet:
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, m_bullet.group(1))
            i += 1
            continue
        if m_num:
            # Explicit numbers: Word's 'List Number' style would continue numbering across lists
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, f'{m_num.group(1)}.\t{m_num.group(2)}')
            i += 1
            continue

        # Figure caption (italic line starting with *Figure / *Table)
        if stripped.startswith('*Figure') or stripped.startswith('*Table'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped, size=Pt(11))
            i += 1
            continue

        # Paragraph (merge consecutive text lines)
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|>|[*\-]\s|\d+\.\s|---|\*\*)', lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if para[0].startswith('**Table'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True   # keep table caption with its table
        add_inline(p, ' '.join(para))

    doc.save(out)
    print(f'Saved: {out}')


if __name__ == '__main__':
    import sys
    # optional: python build_report_docx.py <input.md> <output.docx>
    main(*sys.argv[1:3]) if len(sys.argv) >= 3 else main()
